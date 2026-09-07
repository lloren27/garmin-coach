from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any


def parse_lab_test(path: Path) -> dict[str, Any]:
    text = extract_text(path)
    extracted = extract_metrics(text)
    profile_update = build_profile_update(extracted)
    notes = build_notes(text, extracted)
    return {
        "extracted": extracted,
        "profile_update": profile_update,
        "notes": notes,
        "raw_excerpt": text[:2500],
    }


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix == ".docx":
        return extract_docx_text(path)
    raise ValueError(f"Unsupported lab test document: {path.name}")


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return clean_text("\n".join(parts))


def extract_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return clean_text("\n".join(parts))


def extract_metrics(text: str) -> dict[str, Any]:
    normalized = normalize(text)
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    extracted: dict[str, Any] = {}

    extracted["max_hr"] = find_metric(lines, ("fcmax", "fc maxima", "frecuencia cardiaca maxima", "hr max"), 120, 230)
    extracted["resting_hr"] = find_metric(lines, ("fc reposo", "frecuencia cardiaca reposo", "reposo"), 30, 100)
    extracted["vt1_hr"] = find_metric(lines, ("vt1", "umbral aerobico", "primer umbral"), 70, 190)
    extracted["vt2_hr"] = find_metric(lines, ("vt2", "umbral anaerobico", "segundo umbral"), 90, 220)
    extracted["lactate_hr"] = extracted.get("vt2_hr") or find_metric(
        lines,
        ("umbral lactato", "umbral ventilatorio 2", "fc umbral"),
        90,
        220,
    )
    extracted["vo2max"] = find_metric(lines, ("vo2max", "vo2 max", "consumo maximo de oxigeno"), 25, 90)
    extracted["weight_kg"] = find_metric(lines, ("peso", "weight"), 40, 160)
    extracted["ftp"] = find_power(lines, ("ftp", "potencia funcional"))
    extracted["vt1_power"] = find_power(lines, ("vt1", "umbral aerobico", "primer umbral"))
    extracted["vt2_power"] = find_power(lines, ("vt2", "umbral anaerobico", "segundo umbral"))
    extracted["vt1_pace"] = find_pace(lines, ("vt1", "umbral aerobico", "primer umbral"))
    extracted["vt2_pace"] = find_pace(lines, ("vt2", "umbral anaerobico", "segundo umbral"))
    extracted["threshold_pace"] = extracted.get("vt2_pace") or find_pace(lines, ("ritmo umbral", "umbral lactato"))

    return {key: value for key, value in extracted.items() if value not in (None, "")}


def build_profile_update(extracted: dict[str, Any]) -> dict[str, Any]:
    update: dict[str, Any] = {}
    for source, target in (
        ("max_hr", "max_hr"),
        ("resting_hr", "resting_hr"),
        ("lactate_hr", "lactate_hr"),
        ("vt1_hr", "vt1_hr"),
        ("vt2_hr", "vt2_hr"),
        ("vo2max", "vo2max"),
        ("ftp", "ftp"),
        ("vt1_power", "vt1_power"),
        ("vt2_power", "vt2_power"),
        ("weight_kg", "weight_kg"),
    ):
        if extracted.get(source) is not None:
            update[target] = extracted[source]
    if extracted.get("threshold_pace"):
        update["running_threshold_pace"] = extracted["threshold_pace"]
    return update


def build_notes(text: str, extracted: dict[str, Any]) -> list[str]:
    normalized = normalize(text)
    notes = []
    if not extracted:
        notes.append("No se detectaron metricas estructuradas; revisar el informe manualmente.")
    if any(token in normalized for token in ("no apto", "arritmia", "isquemia", "hipertension", "limitacion", "contraindicacion")):
        notes.append("El informe parece incluir una alerta o limitacion medica: respetar criterio profesional antes de aplicar carga.")
    if "vt1_hr" in extracted and "vt2_hr" in extracted:
        notes.append("VT1 y VT2 detectados: se usaran como base preferente para zonas de FC.")
    elif "max_hr" in extracted:
        notes.append("Sin VT1/VT2 claros: las zonas se estimaran desde FCmax/FC reposo si esta disponible.")
    return notes


def find_metric(lines: list[str], labels: tuple[str, ...], minimum: float, maximum: float) -> float | int | None:
    for line in matching_lines(lines, labels):
        for value in numbers(line):
            if minimum <= value <= maximum:
                return int(value) if value.is_integer() else round(value, 1)
    return None


def find_power(lines: list[str], labels: tuple[str, ...]) -> int | None:
    for line in matching_lines(lines, labels):
        if not any(token in line for token in (" w", "wat", "potencia", "ftp")):
            continue
        for value in numbers(line):
            if 50 <= value <= 600:
                return int(value)
    return None


def find_pace(lines: list[str], labels: tuple[str, ...]) -> str | None:
    for line in matching_lines(lines, labels):
        match = re.search(r"\b([2-9]:[0-5][0-9])\b", line)
        if match:
            return match.group(1)
    return None


def matching_lines(lines: list[str], labels: tuple[str, ...]) -> list[str]:
    result = []
    for index, line in enumerate(lines):
        if any(label in line for label in labels):
            result.append(line)
            if index + 1 < len(lines):
                result.append(lines[index + 1])
    return result


def numbers(text: str) -> list[float]:
    values = []
    for match in re.finditer(r"(?<![:\d])([0-9]{1,3}(?:[,.][0-9]+)?)(?![:\d])", text):
        try:
            values.append(float(match.group(1).replace(",", ".")))
        except ValueError:
            pass
    return values


def normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.lower())
    ascii_text = "".join(char for char in normalized if not unicodedata.combining(char))
    ascii_text = ascii_text.replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", ascii_text)


def clean_text(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", value.replace("\x00", "")).strip()
